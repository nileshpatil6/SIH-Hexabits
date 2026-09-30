"""Build the iTantra SIH 2026 idea deck on top of the official SIH template.

Keeps every template element (SIH logos, titles, team oval, footer bar, pointer headings)
and replaces the placeholder body text with designed content.
"""
import copy
import io
import os

import qrcode
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
ICO = os.path.join(IMG, "icons")

NAVY = RGBColor(0x1F, 0x3B, 0x73)
BLUE = RGBColor(0x00, 0x70, 0xC0)
SAFF = RGBColor(0xE8, 0x71, 0x1A)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x5B, 0x6B, 0x7F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
T_BLUE = RGBColor(0xEA, 0xF2, 0xFB)
T_SAFF = RGBColor(0xFF, 0xF1, 0xE6)
T_GREEN = RGBColor(0xEA, 0xF6, 0xEE)
T_GRAY = RGBColor(0xF4, 0xF6, 0xF9)
LINE = RGBColor(0xD9, 0xE1, 0xEA)

BODY = "Calibri"
HEAD = "Calibri"
TEAM = "Hexabits"


# ----------------------------------------------------------------------------- helpers
def no_shadow(shape):
    shape.shadow.inherit = False


def box(slide, x, y, w, h, fill=None, line=None, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    no_shadow(s)
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=BODY, space_after=0, line_spacing=None, margin=0.0):
    """paras: str | list of paragraphs; a paragraph is str or list of runs (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, opt in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
            if opt.get("link"):
                r.hyperlink.address = opt["link"]
    return tb


def icon(slide, name, color, x, y, d):
    return slide.shapes.add_picture(os.path.join(ICO, f"{name}_{color}.png"), Inches(x), Inches(y), Inches(d), Inches(d))


def icon_circle(slide, name, fill, x, y, d, pad=0.22):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    icon(slide, name, "w", x + d * pad, y + d * pad, d * (1 - 2 * pad))
    return c


def arrow(slide, x, y, w, h, color=MUTED, kind=MSO_SHAPE.RIGHT_ARROW):
    return box(slide, x, y, w, h, fill=color, shape=kind)


def section(slide, x, y, w, label, color=NAVY, size=15, ico=None, ico_color="n"):
    dx = 0
    if ico:
        icon(slide, ico, ico_color, x, y + 0.02, 0.3)
        dx = 0.38
    text(slide, x + dx, y, w - dx, 0.36, label, size=size, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)


def set_title(slide, title, size=28):
    t = slide.shapes.title
    tf = t.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    runs[0].text = title
    runs[0].font.size = Pt(size)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_team_oval(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and "Team" in sh.text_frame.text and sh.shape_type == 1:
            tf = sh.text_frame
            for extra in tf.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            p = tf.paragraphs[0]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            r = p.runs[0] if p.runs else p.add_run()
            r.text = TEAM
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.CENTER
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def body_textbox(slide):
    for sh in slide.shapes:
        if sh.shape_type == 17 and sh.name.startswith("TextBox"):
            return sh
    return None


def qr_png(url):
    q = qrcode.QRCode(border=1, box_size=10)
    q.add_data(url)
    q.make(fit=True)
    im = q.make_image(fill_color=(31, 59, 115), back_color="white")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf


def step_card(slide, x, y, w, h, ico, fill, tint, title, sub, num):
    box(slide, x, y, w, h, fill=tint, radius=0.14)
    d = 0.58
    icon_circle(slide, ico, fill, x + 0.14, y + (h - d) / 2, d)
    text(slide, x + 0.82, y + 0.1, w - 0.9, 0.34, [[(f"{num}  ", {"color": fill, "bold": True}), (title, {"bold": True})]],
         size=12.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.82, y + 0.44, w - 0.9, h - 0.5, sub, size=9.5, color=MUTED)


# ----------------------------------------------------------------------------- slides
def slide1(s):
    # Subtitle placeholder ("TITLE PAGE") becomes the idea name.
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "TITLE PAGE":
            p = next(pp for pp in sh.text_frame.paragraphs if pp.runs and "TITLE" in "".join(r.text for r in pp.runs))
            p.runs[0].text = "iTantra: Voice That Travels as Meaning"
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].font.size = Pt(24)
            p.runs[0].font.color.rgb = NAVY
    tb = body_textbox(s)
    rows = [
        ("Problem Statement ID – ", "SIH26173"),
        ("Problem Statement Title – ", "iTantra: Indian Multilingual TTS & STT Aided Neural Transceiver Radio Access for low bitrate links"),
        ("Organisation – ", "Indian Space Research Organisation (ISRO)"),
        ("Theme – ", "Miscellaneous"),
        ("PS Category – ", "Software"),
        ("Team ID – ", ""),
        ("Team Name (Registered on portal) – ", TEAM),
    ]
    tf = tb.text_frame
    first = tf.paragraphs[0]
    tmpl_ppr = copy.deepcopy(first._p.pPr) if first._p.pPr is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, (label, value) in enumerate(rows):
        p = first if i == 0 else tf.add_paragraph()
        if i and tmpl_ppr is not None:
            p._p.insert(0, copy.deepcopy(tmpl_ppr))
        p.space_after = Pt(9)
        r = p.add_run()
        r.text = label
        r.font.bold = True
        r.font.size = Pt(15)
        r.font.name = BODY
        r.font.color.rgb = NAVY
        if value:
            r2 = p.add_run()
            r2.text = value
            r2.font.bold = False
            r2.font.size = Pt(15)
            r2.font.name = BODY
            r2.font.color.rgb = INK
    tb.top = Inches(2.05)
    tb.height = Inches(4.4)

    badges = [("phone", GREEN, "Working prototype"), ("lang", SAFF, "10 languages, offline"),
              ("github", NAVY, "Fully open source")]
    for i, (ic, col, lab) in enumerate(badges):
        bx = 0.5 + i * 2.35
        box(s, bx, 5.45, 2.22, 0.5, fill=WHITE, line=col, radius=0.5, lw=1.25)
        icon(s, ic, {GREEN: "g", SAFF: "o", NAVY: "n"}[col], bx + 0.13, 5.55, 0.3)
        text(s, bx + 0.5, 5.45, 1.7, 0.5, lab, size=10.5, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)

    # Brand strip along the bottom.
    box(s, 0.36, 6.5, 12.6, 0.82, fill=T_BLUE, radius=0.5)
    s.shapes.add_picture(os.path.join(IMG, "logo-itantra-v1-20260930-170705-transparent.png"),
                         Inches(0.55), Inches(6.5), Inches(0.82), Inches(0.82))
    text(s, 1.5, 6.5, 11.3, 0.82,
         [[("iTantra  ", {"bold": True, "color": NAVY, "size": 16}),
           ("Offline speech-to-speech radio for disasters  •  10 Indian languages  •  ~234 bytes per spoken sentence  •  "
            "runs on 3 GB Android phones", {"color": INK, "size": 13})]],
         anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "iTantra turns speech into a tiny signed text packet on the sender's phone, relays it phone to phone over "
        "Bluetooth and Wi-Fi Direct, and speaks it again on the receiver. No tower, internet or server.")


def slide2(s):
    set_title(s, "iTantra: Send the Meaning, Not the Audio", 28)
    set_team_oval(s)
    remove_shape(body_textbox(s))

    text(s, 0.45, 1.25, 12.4, 0.4,
         [[("❖  ", {"color": SAFF}), ("Proposed Solution (Describe your Idea/Solution/Prototype)", {})]],
         size=17, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.45, 1.62, 12.4, 0.3,
         "An Android walkie-talkie that turns speech into a few hundred bytes of signed text, relays it phone to phone, "
         "and speaks it again on the other side.",
         size=11.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

    steps = [
        ("mic", SAFF, T_SAFF, "Speak", "Any of 10 Indian languages, push-to-talk or hands-free"),
        ("stt", NAVY, T_BLUE, "On-device STT", "Silero VAD finds the pause; IndicConformer writes the sentence"),
        ("packet", GREEN, T_GREEN, "234 B packet", "Text + language + Ed25519 signature, TTL 7"),
        ("mesh", BLUE, T_BLUE, "Phone mesh", "BLE + Wi-Fi Direct relays, up to 7 hops, no tower"),
        ("speaker", SAFF, T_SAFF, "Spoken aloud", "On-device TTS voice; alerts ring at max volume"),
    ]
    gap, y, h = 0.26, 2.02, 1.02
    w = (12.43 - gap * 4) / 5
    for i, (ic, col, tint, t, sub) in enumerate(steps):
        x = 0.45 + i * (w + gap)
        step_card(s, x, y, w, h, ic, col, tint, t, sub, i + 1)
        if i < 4:
            arrow(s, x + w + 0.04, y + h / 2 - 0.1, gap - 0.08, 0.2, color=RGBColor(0xB8, 0xC4, 0xD2), kind=MSO_SHAPE.CHEVRON)

    # Hero image + stat chips
    s.shapes.add_picture(os.path.join(IMG, "hero_w.png"), Inches(0.5), Inches(3.18), Inches(5.6), Inches(3.13))
    chips = [("≈234 B", "per spoken sentence"), ("30× smaller", "than Opus 16 kbps"), ("120× smaller", "than 64 kbps voice")]
    cw = (5.75 - 0.2) / 3
    for i, (big, small) in enumerate(chips):
        x = 0.45 + i * (cw + 0.1)
        box(s, x, 6.33, cw, 0.54, fill=NAVY, radius=0.3)
        text(s, x, 6.33, cw, 0.54, [[(big, {"bold": True, "color": WHITE, "size": 12.5})],
                                     [(small, {"color": RGBColor(0xDC, 0xE6, 0xF5), "size": 9})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # How it addresses the problem
    rx, rw = 6.5, 6.38
    section(s, rx, 3.18, rw, "How it addresses the problem", ico="check", ico_color="g")
    bullets = [
        ("Low bitrate: ", "one sentence is ≈234 B (≈0.5 kbps) instead of >64 kbps streamed audio"),
        ("No network: ", "phones relay for each other; no tower, internet, SIM or server"),
        ("Inclusive: ", "people who cannot read still hear the message, in their own language"),
        ("Low-end ready: ", "int8 models, one language in RAM, runs on 3 GB Android phones"),
    ]
    for i, (b, t) in enumerate(bullets):
        yy = 3.58 + i * 0.36
        icon(s, "check", "g", rx + 0.04, yy + 0.06, 0.22)
        text(s, rx + 0.36, yy, rw - 0.36, 0.34, [[(b, {"bold": True, "color": NAVY}), (t, {})]],
             size=11.5, anchor=MSO_ANCHOR.MIDDLE)

    section(s, rx, 5.06, rw, "Innovation and uniqueness", ico="rocket", ico_color="o")
    cards = [
        ("wave", SAFF, "Semantic radio", "Sends meaning + prosody, not the waveform"),
        ("alert", SAFF, "Alert-first mesh", "Floods all hops; loud until acknowledged"),
        ("lock", NAVY, "Tamper-proof", "Signed packets; encrypted private messages"),
        ("offline", GREEN, "100% on-device", "No cloud API; language packs download once"),
    ]
    cw2, ch = (rw - 0.12) / 2, 0.66
    for i, (ic, col, t, sub) in enumerate(cards):
        cx = rx + (i % 2) * (cw2 + 0.12)
        cy = 5.47 + (i // 2) * (ch + 0.08)
        box(s, cx, cy, cw2, ch, fill=T_GRAY, line=LINE, radius=0.16)
        icon_circle(s, ic, col, cx + 0.1, cy + 0.12, 0.42)
        text(s, cx + 0.62, cy + 0.06, cw2 - 0.7, 0.28, t, size=11.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + 0.62, cy + 0.33, cw2 - 0.7, 0.3, sub, size=9.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Pipeline: speak, on-device STT, a ~234 byte signed packet, phone-to-phone mesh, on-device TTS. "
        "The byte figures are measured from our packet format: 22 B header, text, 64 B signature.")


def slide3(s):
    set_title(s, "TECHNICAL APPROACH", 28)
    set_team_oval(s)
    remove_shape(body_textbox(s))

    lx, lw = 0.45, 8.15
    section(s, lx, 1.25, lw, "Technologies to be used", ico="code", ico_color="n")
    rows = [
        ("App & UI", NAVY, ["Flutter (Dart)", "Kotlin (native Android)", "Pigeon typed bridge", "Riverpod"]),
        ("On-device AI", SAFF, ["sherpa-onnx runtime", "IndicConformer int8 (STT)", "MMS VITS (TTS)", "Silero VAD"]),
        ("Mesh & security", GREEN, ["BLE GATT mesh", "Wi-Fi Direct", "Ed25519 signatures", "X25519 + ChaCha20"]),
        ("Model tooling", BLUE, ["Python", "ONNX int8 quantisation", "Google Colab", "Hugging Face model packs"]),
    ]
    tints = {NAVY: T_BLUE, SAFF: T_SAFF, GREEN: T_GREEN, BLUE: T_BLUE}
    for i, (cat, col, chips) in enumerate(rows):
        y = 1.68 + i * 0.44
        box(s, lx, y, 1.55, 0.36, fill=col, radius=0.5)
        text(s, lx, y, 1.55, 0.36, cat, size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x = lx + 1.68
        for c in chips:
            cw = 0.066 * len(c) + 0.32
            box(s, x, y, cw, 0.36, fill=tints[col], radius=0.5)
            text(s, x, y, cw, 0.36, c, size=10, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            x += cw + 0.1

    section(s, lx, 3.52, lw, "Methodology and process for implementation (flow)", ico="route", ico_color="o")
    bw, gap = 1.82, 0.29
    sender = [
        ("mic", "Mic capture", "16 kHz mono, echo + noise cancel"),
        ("wave", "Silero VAD", "Pause = end of sentence"),
        ("stt", "IndicConformer STT", "int8 CTC, one language resident"),
        ("packet", "Packetiser", "Sign, fragment, TTL 7"),
    ]
    receiver = [
        ("shield", "Verify + reassemble", "Drop unsigned or replayed"),
        ("speaker", "MMS VITS TTS", "Voice in sender's language"),
        ("hearing", "Voice note", "Text shown with audio"),
        ("alert", "ALERT mode", "Alarm stream, max volume, x3"),
    ]

    def lane(items, y, col, tint, label):
        text(s, lx, y - 0.02, 1.2, 0.2, label, size=8.5, bold=True, color=col)
        for i, (ic, t, sub) in enumerate(items):
            x = lx + i * (bw + gap)
            box(s, x, y + 0.18, bw, 0.76, fill=tint, line=None, radius=0.14)
            icon(s, ic, "n" if col == NAVY else ("o" if col == SAFF else "g"), x + 0.1, y + 0.3, 0.3)
            text(s, x + 0.45, y + 0.22, bw - 0.5, 0.3, t, size=10.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.12, y + 0.54, bw - 0.18, 0.38, sub, size=9, color=MUTED)
            if i < 3:
                arrow(s, x + bw + 0.05, y + 0.48, gap - 0.1, 0.16, color=RGBColor(0xB8, 0xC4, 0xD2))

    lane(sender, 3.9, NAVY, T_BLUE, "PHONE A  ·  SENDER")
    # Mesh band
    my = 5.04
    arrow(s, lx + 3 * (bw + gap) + bw / 2 - 0.1, 4.86, 0.2, 0.2, color=RGBColor(0xB8, 0xC4, 0xD2), kind=MSO_SHAPE.DOWN_ARROW)
    box(s, lx, my + 0.02, lw, 0.46, fill=NAVY, radius=0.5)
    icon(s, "bt", "w", lx + 0.18, my + 0.1, 0.3)
    icon(s, "wifi", "w", lx + 0.52, my + 0.1, 0.3)
    text(s, lx + 0.95, my + 0.02, lw - 1.1, 0.46,
         [[("Mesh:  ", {"bold": True, "color": WHITE}),
           ("verify → dedup → relay with jitter (≤7 hops) → store-and-forward outbox when no neighbour",
            {"color": RGBColor(0xDC, 0xE6, 0xF5)})]],
         size=10.5, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, lx + bw - 0.4, my + 0.5, 0.2, 0.2, color=RGBColor(0xB8, 0xC4, 0xD2), kind=MSO_SHAPE.DOWN_ARROW)
    lane(receiver, 5.66, GREEN, T_GREEN, "PHONE B  ·  RECEIVER")

    # Prototype panel
    px = 8.92
    section(s, px, 1.25, 3.95, "Working prototype (Android)", ico="phone", ico_color="g")
    # phone mockup, recreated from the real app's Talk screen
    fx, fy, fw, fh = 9.0, 1.7, 2.1, 4.45
    box(s, fx, fy, fw, fh, fill=RGBColor(0x1F, 0x29, 0x37), radius=0.12)
    sx, sy, sw, sh = fx + 0.09, fy + 0.12, fw - 0.18, fh - 0.24
    box(s, sx, sy, sw, sh, fill=RGBColor(0xF3, 0xF8, 0xF4), radius=0.08)
    text(s, sx + 0.1, sy + 0.08, sw - 0.2, 0.3,
         [[("Talk · ", {"color": INK}), ("हिन्दी", {"font": "Nirmala UI", "color": INK})]], size=12, bold=True)
    box(s, sx, sy + 0.44, sw, 0.26, fill=RGBColor(0xE2, 0xE8, 0xE4), radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    text(s, sx + 0.05, sy + 0.44, sw - 0.1, 0.26,
         [[("● ", {"color": GREEN}), ("STT ready   ", {}), ("● ", {"color": GREEN}), ("Voice   ", {}), ("● ", {"color": GREEN}), ("Offline", {})]],
         size=7, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    bx = sx + 0.5
    box(s, bx, sy + 0.9, sw - 0.6, 0.7, fill=RGBColor(0xA7, 0xF0, 0xCF), radius=0.3)
    text(s, bx + 0.1, sy + 0.93, sw - 0.8, 0.36, "हलो हलो", size=15, color=RGBColor(0x0B, 0x4F, 0x3A), font="Nirmala UI",
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, bx + 0.1, sy + 1.28, sw - 0.8, 0.26, "Hindi · STT 236 ms", size=7.5, color=RGBColor(0x2F, 0x6B, 0x55),
         anchor=MSO_ANCHOR.MIDDLE)
    box(s, sx + 0.1, sy + sh - 1.2, sw - 0.2, 0.3, fill=WHITE, line=LINE, radius=0.2)
    text(s, sx + 0.18, sy + sh - 1.2, sw - 0.4, 0.3, "Or type a message", size=7.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    icon_circle(s, "mic", RGBColor(0x1D, 0x6B, 0x4F), sx + sw - 0.72, sy + sh - 0.8, 0.62, pad=0.26)
    box(s, sx + 0.1, sy + sh - 0.62, 0.8, 0.26, fill=RGBColor(0xD6, 0xEA, 0xDD), radius=0.5)
    text(s, sx + 0.1, sy + sh - 0.62, 0.8, 0.26, "PTT | Call", size=7.5, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    metrics = [("236 ms", "Hindi sentence decode", SAFF), ("644 ms", "STT model load", NAVY),
               ("654 ms", "TTS voice load", NAVY), ("10 / 10", "languages verified", GREEN)]
    for i, (big, lab, col) in enumerate(metrics):
        y = 1.7 + i * 1.13
        box(s, 11.3, y, 1.58, 1.03, fill=T_GRAY, line=LINE, radius=0.14)
        text(s, 11.3, y + 0.08, 1.58, 0.5, big, size=20, bold=True, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 11.35, y + 0.56, 1.48, 0.42, lab, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 8.95, 6.25, 3.95, 0.55,
         "Talk screen of the running app (redrawn); timings measured on a OnePlus CPH2717, Android 16.",
         size=9, color=MUTED, align=PP_ALIGN.LEFT)
    s.notes_slide.notes_text_frame.text = (
        "Everything runs on the phone: Flutter UI, Kotlin native layer, sherpa-onnx for STT, TTS and VAD. "
        "The prototype is running on a OnePlus today; timings on the right are measured on that phone.")


def slide4(s):
    set_title(s, "FEASIBILITY AND VIABILITY", 28)
    set_team_oval(s)
    remove_shape(body_textbox(s))

    lx, lw = 0.45, 5.8
    section(s, lx, 1.25, lw, "Analysis of the feasibility of the idea", ico="science", ico_color="n")
    tiles = [("10 / 10", "languages running fully offline", GREEN, T_GREEN),
             ("~0.3 GB", "per language (speech + voice packs)", NAVY, T_BLUE),
             ("<1.1 GB", "peak RAM (est.), fits 3 GB phones", SAFF, T_SAFF)]
    tw = (lw - 0.2) / 3
    for i, (big, lab, col, tint) in enumerate(tiles):
        x = lx + i * (tw + 0.1)
        box(s, x, 1.68, tw, 1.0, fill=tint, radius=0.14)
        text(s, x, 1.72, tw, 0.48, big, size=22, bold=True, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.08, 2.2, tw - 0.16, 0.44, lab, size=9.5, color=INK, align=PP_ALIGN.CENTER)

    cd = CategoryChartData()
    langs = ["Kannada", "Marathi", "Tamil", "Bengali", "Gujarati", "Telugu", "Odia", "Hindi", "English", "Malayalam"]
    cer = [1.0, 1.2, 2.0, 3.2, 3.3, 3.3, 6.0, 8.5, 10.2, 13.6]
    cd.categories = langs
    cd.add_series("Character error %", cer)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(lx), Inches(2.82), Inches(lw), Inches(3.3), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Round-trip check: character error % by language (lower is better)"
    tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
    tr.font.size = Pt(11)
    tr.font.bold = True
    tr.font.color.rgb = NAVY
    tr.font.name = BODY
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(9)
    dl.font.color.rgb = INK
    ser = plot.series[0]
    for i, v in enumerate(cer):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = GREEN if v <= 5 else (NAVY if v <= 10 else SAFF)
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(8.5)
    ca.tick_labels.font.color.rgb = MUTED
    ca.format.line.color.rgb = LINE
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xEE, 0xF1, 0xF5)
    va.maximum_scale = 16
    text(s, lx, 6.16, lw, 0.66,
         "Each language: our TTS voice speaks a flood warning, our STT transcribes it back (sherpa-onnx, CPU, 2 threads). "
         "STT ran 5–20× faster than real time. Next: real speech from 10 speakers per language.",
         size=9, color=MUTED)

    rx, rw = 6.55, 6.33
    section(s, rx, 1.25, rw, "Potential challenges and risks  →  Strategies for overcoming them", ico="warning", ico_color="o",
            size=14)
    risks = [
        ("route", "Short range per hop", "~60–100 m per Bluetooth hop",
         "7-hop relay, BLE Coded PHY where supported, store-and-forward by moving phones"),
        ("hearing", "Hard speech", "Noise, soft voices, Malayalam / Odia accuracy",
         "Gain control + tuned VAD, distress hotword list, text shown with voice, real-speech test set"),
        ("battery", "Low-end phones", "RAM, battery drain, background kills",
         "int8 models, one language in RAM, duty-cycled scanning, foreground service"),
        ("shield", "Fake alerts and misuse", "Spoofing; mesh apps were flagged in India (Jul 2026)",
         "Ed25519-signed packets, verified NDMA / ISRO key, built as an authenticated alert tool"),
        ("book", "Voice licence", "MMS-TTS voices are non-commercial",
         "Swap to CC-BY voices trained on SYSPIN / IndicVoices-R before deployment"),
    ]
    rh, rg = 0.96, 0.07
    lw2, aw = 2.72, 0.3
    for i, (ic, t, sub, fix) in enumerate(risks):
        y = 1.68 + i * (rh + rg)
        box(s, rx, y, lw2, rh, fill=T_SAFF, radius=0.14)
        icon(s, ic, "o", rx + 0.12, y + 0.14, 0.32)
        text(s, rx + 0.52, y + 0.1, lw2 - 0.6, 0.36, t, size=11.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, rx + 0.14, y + 0.48, lw2 - 0.24, 0.46, sub, size=9.5, color=MUTED)
        arrow(s, rx + lw2 + 0.05, y + rh / 2 - 0.1, aw - 0.1, 0.2, color=RGBColor(0xB8, 0xC4, 0xD2), kind=MSO_SHAPE.CHEVRON)
        mx = rx + lw2 + aw
        box(s, mx, y, rw - lw2 - aw, rh, fill=T_GREEN, radius=0.14)
        icon(s, "check", "g", mx + 0.12, y + (rh - 0.3) / 2, 0.3)
        text(s, mx + 0.52, y + 0.06, rw - lw2 - aw - 0.62, rh - 0.12, fix, size=10.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Feasibility is proven, not assumed: the app runs today and all 10 languages pass a round-trip check. "
        "RAM figure is an engineering estimate to be replaced with dumpsys meminfo on a 3 GB phone.")


def slide5(s):
    set_title(s, "IMPACT AND BENEFITS", 28)
    set_team_oval(s)
    remove_shape(body_textbox(s))

    stats = [
        ("85%", "of Indians speak one of our 10 languages as mother tongue (Census 2011)", SAFF),
        ("~26%", "of Indians aged 7+ cannot read; voice still reaches them (Census 2011)", NAVY),
        ("0", "towers, internet, SIM cards or servers needed", GREEN),
        ("₹0", "data cost per message, on phones people already own", BLUE),
    ]
    sw = (12.43 - 0.3) / 4
    for i, (big, lab, col) in enumerate(stats):
        x = 0.45 + i * (sw + 0.1)
        box(s, x, 1.28, sw, 1.08, fill=T_GRAY, line=LINE, radius=0.14)
        nw = 0.26 * len(big) + 0.35
        text(s, x + 0.18, 1.3, nw, 1.04, big, size=30, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.2 + nw, 1.3, sw - nw - 0.32, 1.04, lab, size=10.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    section(s, 0.45, 2.5, 5.9, "Potential impact on the target audience", ico="groups", ico_color="o")
    s.shapes.add_picture(os.path.join(IMG, "impact_w.png"), Inches(0.45), Inches(2.9), Inches(5.85), Inches(3.92))

    rx, rw = 6.55, 6.33
    section(s, rx, 2.5, rw, "Benefits of the solution", ico="handshake", ico_color="g")
    cards = [
        ("access", SAFF, T_SAFF, "Social", "Voice in 10 languages for non-literate, elderly and visually impaired people; no typing needed."),
        ("rupee", GREEN, T_GREEN, "Economic", "Runs on existing low-cost phones; zero infrastructure, data or SMS cost; open models."),
        ("eco", GREEN, T_GREEN, "Environmental", "No diesel towers or generators; low-power Bluetooth saves battery in long outages."),
        ("satellite", NAVY, T_BLUE, "Strategic (ISRO)", "Byte-sized frames suit NavIC messaging and LoRa links; data never leaves Indian devices."),
    ]
    cw, chh = (rw - 0.12) / 2, 1.42
    for i, (ic, col, tint, t, sub) in enumerate(cards):
        cx = rx + (i % 2) * (cw + 0.12)
        cy = 2.9 + (i // 2) * (chh + 0.1)
        box(s, cx, cy, cw, chh, fill=tint, radius=0.12)
        icon_circle(s, ic, col, cx + 0.14, cy + 0.14, 0.5)
        text(s, cx + 0.76, cy + 0.14, cw - 0.86, 0.5, t, size=13, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + 0.16, cy + 0.72, cw - 0.3, chh - 0.78, sub, size=10.5, color=INK)
    fy = 2.9 + 2 * (chh + 0.1)
    box(s, rx, fy, rw, 6.82 - fy, fill=NAVY, radius=0.2)
    icon(s, "rocket", "w", rx + 0.16, fy + (6.82 - fy - 0.34) / 2, 0.34)
    text(s, rx + 0.62, fy, rw - 0.74, 6.82 - fy,
         [[("Next:  ", {"bold": True, "color": WHITE}),
           ("\"heard by N phones\" receipts  •  SMS gateway when any phone regains signal  •  voice-preserving prosody  •  "
            "10-phone field trial", {"color": RGBColor(0xDC, 0xE6, 0xF5)})]],
         size=10.5, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Census 2011: Hindi, Bengali, Marathi, Telugu, Tamil, Gujarati, Kannada, Odia and Malayalam together are ~85% "
        "of mother-tongue speakers; literacy 74% means ~26% of people aged 7+ cannot read.")


def slide6(s):
    set_title(s, "RESEARCH AND REFERENCES", 28)
    set_team_oval(s)
    remove_shape(body_textbox(s))
    section(s, 0.45, 1.25, 12.4, "Details / Links of the reference and research work", ico="book", ico_color="n")

    cols = [
        ("layers", NAVY, T_BLUE, "Models and tools", [
            ("AI4Bharat IndicConformer ASR (MIT)", "github.com/AI4Bharat/IndicConformerASR"),
            ("sherpa-onnx: offline speech runtime", "github.com/k2-fsa/sherpa-onnx"),
            ("Meta MMS-TTS voices (VITS)", "huggingface.co/facebook/mms-tts"),
            ("Silero VAD", "github.com/snakers4/silero-vad"),
            ("IndicConformer sherpa-onnx int8 exports", "huggingface.co/parismitaglobalsolutions/indicconformer-sherpa-onnx"),
        ]),
        ("science", SAFF, T_SAFF, "Research", [
            ("IndicVoices: multilingual Indian speech (ACL 2024)", "arxiv.org/abs/2403.01926"),
            ("STCTS: speech at 80 bps via text + prosody", "arxiv.org/abs/2512.00451"),
            ("IndicVoices-R: Indian TTS corpus (NeurIPS 2024)", "github.com/AI4Bharat/IndicVoices-R"),
            ("BitChat mesh protocol whitepaper", "github.com/permissionlesstech/bitchat"),
            ("SYSPIN: open CC-BY Indian TTS dataset (IISc)", "vaani.iisc.ac.in/dataset/syspindataset"),
        ]),
        ("public", GREEN, T_GREEN, "Real-world evidence", [
            ("FireChat mesh, Hong Kong 2014 (500k installs)", "en.wikipedia.org/wiki/FireChat"),
            ("BitChat in Nepal 2025: 48,781 installs in a day", "forbes.com (Sep 11, 2025)"),
            ("Census of India 2011: languages, literacy", "censusindia.gov.in"),
            ("NDMA: disaster management in India", "ndma.gov.in"),
            ("Bridgefy offline mesh adds E2E encryption (2020)", "techcrunch.com (Nov 2, 2020)"),
        ]),
    ]
    links = {
        "techcrunch.com (Nov 2, 2020)": "https://techcrunch.com/2020/11/02/bridgefy-launches-end-to-end-encrypted-messaging-for-the-app-used-during-protests-and-disasters/",
        "forbes.com (Sep 11, 2025)": "https://www.forbes.com/sites/digital-assets/2025/09/11/jack-dorseys-bitchat-gains-traction-during-nepals-unrest/",
    }
    cw = (12.43 - 0.3) / 3
    for i, (ic, col, tint, t, refs) in enumerate(cols):
        x = 0.45 + i * (cw + 0.15)
        box(s, x, 1.7, cw, 3.72, fill=tint, radius=0.08)
        icon_circle(s, ic, col, x + 0.2, 1.84, 0.5)
        text(s, x + 0.82, 1.84, cw - 1.0, 0.5, t, size=14, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        for j, (name, url) in enumerate(refs):
            y = 2.45 + j * 0.6
            href = links.get(url, "https://" + url)
            shown = url if len(url) < 48 else url.split("/")[0] + "/.../" + url.split("/")[-1]
            text(s, x + 0.22, y, cw - 0.4, 0.58,
                 [[(name, {"bold": True, "color": INK, "size": 10.5})],
                  [(shown, {"color": BLUE, "size": 9, "link": href})]])

    # Our work + QR codes
    y0 = 5.58
    box(s, 0.45, y0, 12.43, 1.25, fill=T_GRAY, line=LINE, radius=0.14)
    s.shapes.add_picture(os.path.join(IMG, "logo-itantra-v1-20260930-170705-transparent.png"),
                         Inches(0.6), Inches(y0 + 0.07), Inches(1.1), Inches(1.1))
    text(s, 1.85, y0 + 0.14, 5.0, 1.0,
         [[("Our work is open source", {"bold": True, "color": NAVY, "size": 14})],
          [("Android app, mesh protocol spec and model tooling, plus 20 ready language packs "
            "(10 STT + 10 TTS) verified end to end.", {"color": INK, "size": 10.5})]])
    qrs = [("App source code", "github.com/nileshpatil6/SIH-Hexabits", ["github.com/nileshpatil6/", "SIH-Hexabits"]),
           ("Offline model packs", "huggingface.co/datasets/Mr66/itantra-packs", ["huggingface.co/datasets/", "Mr66/itantra-packs"])]
    for k, (lab, url, lines) in enumerate(qrs):
        qx = 6.75 + k * 3.05
        s.shapes.add_picture(qr_png("https://" + url), Inches(qx), Inches(y0 + 0.1), Inches(1.05), Inches(1.05))
        text(s, qx + 1.15, y0 + 0.16, 1.85, 1.0,
             [[(lab, {"bold": True, "color": NAVY, "size": 11})]] +
             [[(ln, {"color": BLUE, "size": 9, "link": "https://" + url})] for ln in lines])
    s.notes_slide.notes_text_frame.text = "All links are clickable in the PDF export."


def drop_slide(prs, index):
    sld = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


def main():
    prs = Presentation(os.path.join(HERE, "template.pptx"))
    drop_slide(prs, 6)  # "Important instructions" slide is not submitted
    for fn, sl in zip([slide1, slide2, slide3, slide4, slide5, slide6], prs.slides):
        fn(sl)
    out = os.path.join(HERE, "iTantra_SIH2026_Hexabits.pptx")
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
