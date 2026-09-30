"""iTantra SIH 2026 idea deck, v2 "Aero-Tactile Minimal" redesign.

Built on the official SIH template: keeps SIH logos, slide titles, team oval, footer bar and
pointer headings. Adds a light textured background, elevated white cards with soft layered
shadows, gradient icon badges and a repeated radio-pulse ring motif.
"""
import copy
import io
import os

import qrcode
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
ICO = os.path.join(IMG, "icons")

NAVY = RGBColor(0x0F, 0x27, 0x44)
NAVY2 = RGBColor(0x1E, 0x4A, 0x8A)
SAFF = RGBColor(0xF9, 0x73, 0x16)
SAFF2 = RGBColor(0xFB, 0xA5, 0x4A)
EMER = RGBColor(0x10, 0xB9, 0x81)
EMER2 = RGBColor(0x0E, 0x9F, 0x6E)
SKY = RGBColor(0x02, 0x84, 0xC7)
SKY2 = RGBColor(0x38, 0xBD, 0xF8)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
HAIR = RGBColor(0xE2, 0xE8, 0xF0)
ICE = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
T_SAFF = RGBColor(0xFF, 0xF4, 0xEB)
T_EMER = RGBColor(0xEC, 0xFD, 0xF5)
T_SKY = RGBColor(0xEF, 0xF6, 0xFF)
T_SLATE = RGBColor(0xF1, 0xF5, 0xF9)

GRAD = {  # badge gradients (start, end)
    "saff": (SAFF2, SAFF),
    "emer": (RGBColor(0x34, 0xD3, 0x99), EMER2),
    "navy": (NAVY2, NAVY),
    "sky": (SKY2, SKY),
}
BODY = "Calibri"
TEAM = "Hexabits"
LOGO = os.path.join(IMG, "logo-itantra-v1-20260930-170705-transparent.png")


# ----------------------------------------------------------------------------- low-level helpers
def _effects(shape, xml):
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    spPr.append(etree.fromstring(
        f'<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">{xml}</a:effectLst>'))


def soft_shadow(shape, blur=18, dist=4, alpha=9):
    _effects(shape, f'<a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}" dir="5400000" '
                    f'algn="t" rotWithShape="0"><a:srgbClr val="0F172A"><a:alpha val="{alpha * 1000}"/></a:srgbClr>'
                    f'</a:outerShdw>')


def no_effects(shape):
    _effects(shape, "")


def line_alpha(shape, pct):
    ln = shape._element.spPr.find(qn("a:ln"))
    clr = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(pct * 1000)))


def fill_alpha(shape, pct):
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(pct * 1000)))


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=0.75, radius=None, shadow=False, grad=None, angle=90):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if grad:
        s.fill.gradient()
        s.fill.gradient_angle = angle
        st = s.fill.gradient_stops
        st[0].color.rgb, st[0].position = grad[0], 0
        st[1].color.rgb, st[1].position = grad[1], 1.0
    elif fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if shadow:
        soft_shadow(s)
    else:
        no_effects(s)
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, fill=WHITE, radius=0.08, line=HAIR, shadow=True):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, lw=0.6, radius=radius, shadow=shadow)


def text(slide, x, y, w, h, paras, size=11, color=SLATE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=BODY, space_after=0, margin=0.0, line_spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, opt in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
            if opt.get("link"):
                r.hyperlink.address = opt["link"]
    return tb


def pic(slide, path, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def icon(slide, name, color, x, y, d):
    return pic(slide, os.path.join(ICO, f"{name}_{color}.png"), x, y, d, d)


def badge(slide, name, g, x, y, d, pad=0.24, glow=False):
    """Gradient circle with a white icon: the deck's icon treatment."""
    c = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, grad=GRAD[g], angle=45)
    if glow:
        _effects(c, f'<a:outerShdw blurRad="{int(10 * 12700)}" dist="{int(2 * 12700)}" dir="5400000" algn="t" '
                    f'rotWithShape="0"><a:srgbClr val="{str(GRAD[g][1])}"><a:alpha val="30000"/></a:srgbClr></a:outerShdw>')
    icon(slide, name, "w", x + d * pad, y + d * pad, d * (1 - 2 * pad))
    return c


def pulse(slide, cx, cy, r0, color=SAFF, n=3, step=0.22, alpha=(45, 28, 14), lw=1.0):
    """Signature motif: concentric hairline radio rings."""
    for i in range(n):
        r = r0 + step * i
        s = shape(slide, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, line=color, lw=lw)
        line_alpha(s, alpha[min(i, len(alpha) - 1)])


def chip(slide, x, y, label, fill=WHITE, color=SLATE, line=HAIR, size=9.5, h=0.3, bold=False, pad=0.28):
    w = 0.062 * len(label) * (size / 9.5) + pad
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, lw=0.6, radius=0.5)
    text(slide, x, y, w, h, label, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def connector(slide, x1, y1, x2, y2, color=RGBColor(0xB6, 0xC2, 0xD1), lw=1.25, dash=False, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "sysDot")
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "sm")
        t.set("h", "sm")
    return c


def heading(slide, x, y, w, num, label, color=NAVY, size=14.5):
    """Pointer heading with a numbered saffron tag."""
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y + 0.04, 0.42, 0.3, grad=GRAD["saff"], radius=0.5, angle=0)
    text(slide, x, y + 0.04, 0.42, 0.3, num, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.52, y, w - 0.52, 0.38, label, size=size, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)


def to_back(slide, shp):
    tree = slide.shapes._spTree
    el = shp._element
    tree.remove(el)
    tree.insert(2, el)


def background(slide):
    bg = pic(slide, os.path.join(IMG, "v2_bg.png"), 0, 0, 13.333, 7.5)
    to_back(slide, bg)


def set_title(slide, title, size=26):
    t = slide.shapes.title
    tf = t.text_frame
    p = next(pp for pp in tf.paragraphs if pp.runs)
    runs = p.runs
    runs[0].text = title
    runs[0].font.size = Pt(size)
    runs[0].font.color.rgb = NAVY
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in list(tf.paragraphs):
        if extra._p is not p._p:
            extra._p.getparent().remove(extra._p)


def set_team_oval(slide):
    for sh in slide.shapes:
        if sh.shape_type == 1 and sh.has_text_frame and "Team" in sh.text_frame.text:
            tf = sh.text_frame
            for extra in tf.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            p = tf.paragraphs[0]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            r = p.runs[0] if p.runs else p.add_run()
            r.text = TEAM
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.CENTER
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def remove(shp):
    shp._element.getparent().remove(shp._element)


def body_textbox(slide):
    return next(sh for sh in slide.shapes if sh.shape_type == 17 and sh.name.startswith("TextBox"))


def qr_png(url):
    q = qrcode.QRCode(border=1, box_size=10)
    q.add_data(url)
    q.make(fit=True)
    im = q.make_image(fill_color=(15, 39, 68), back_color="white")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf


def content_slide(s, title):
    background(s)
    set_title(s, title)
    set_team_oval(s)
    remove(body_textbox(s))


# ----------------------------------------------------------------------------- slide 1
def slide1(s):
    background(s)
    for sh in list(s.shapes):
        if sh.name in ("Picture 4", "Freeform: Shape 26", "TextBox 9", "Subtitle 3"):
            remove(sh)
    t = s.shapes.title
    for r in t.text_frame.paragraphs[0].runs:
        r.font.color.rgb = NAVY

    # hero render + pulse rings behind it
    pulse(s, 9.75, 4.05, 1.9, color=SAFF, n=3, step=0.42, alpha=(22, 14, 8), lw=1.0)
    pic(s, os.path.join(IMG, "v2_title_hero_c.png"), 7.1, 1.55, w=5.4)

    x0 = 0.55
    chip(s, x0, 1.55, "SIH 2026  •  PS SIH26173  •  ISRO", fill=T_SKY, color=NAVY2, line=None, size=10, h=0.34, bold=True)
    pic(s, LOGO, x0 - 0.05, 2.02, 1.05, 1.05)
    text(s, x0 + 1.05, 1.98, 5.5, 0.8, "iTantra", size=48, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0 + 1.08, 2.72, 5.6, 0.4, [[("Voice that travels as ", {"color": SLATE}), ("meaning", {"color": SAFF, "bold": True})]],
         size=19, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0, 3.25, 6.3, 0.62,
         "Speech in, a 234-byte signed packet across a phone-to-phone mesh, speech out. "
         "No tower, no internet, 10 Indian languages, on 3 GB phones.", size=12.5, color=MUTED)

    # metadata tiles (template pointers kept as labels)
    def tile(x, y, w, h, label, value, vsize=11.5):
        card(s, x, y, w, h, radius=0.14)
        text(s, x + 0.16, y + 0.08, w - 0.3, 0.24, label.upper(), size=8, bold=True, color=MUTED)
        text(s, x + 0.16, y + 0.3, w - 0.3, h - 0.34, value, size=vsize, bold=True, color=NAVY)

    tw = 6.3
    tile(x0, 4.02, tw, 0.8, "Problem Statement Title",
         "iTantra: Indian Multilingual TTS & STT Aided Neural Transceiver Radio Access for low bitrate links", 11)
    c3 = (tw - 0.2) / 3
    tile(x0, 4.92, c3, 0.66, "Problem Statement ID", "SIH26173")
    tile(x0 + c3 + 0.1, 4.92, c3, 0.66, "Theme", "Miscellaneous")
    tile(x0 + 2 * (c3 + 0.1), 4.92, c3, 0.66, "PS Category", "Software")
    tile(x0, 5.68, c3, 0.66, "Team ID", "")
    tile(x0 + c3 + 0.1, 5.68, 2 * c3 + 0.1, 0.66, "Team Name (Registered on portal)", TEAM)

    pills = [("offline", "emer", "100% offline"), ("packet", "saff", "234 B / sentence"),
             ("lang", "sky", "10 languages"), ("sos", "navy", "Alert mode")]
    px = 7.0
    for ic, g, lab in pills:
        w = 0.064 * len(lab) + 0.6
        card(s, px, 6.45, w, 0.42, radius=0.5)
        badge(s, ic, g, px + 0.06, 6.49, 0.34, pad=0.22)
        text(s, px + 0.45, 6.45, w - 0.5, 0.42, lab, size=10.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        px += w + 0.1
    s.notes_slide.notes_text_frame.text = (
        "iTantra turns speech into a tiny signed text packet on the sender's phone, relays it phone to phone over "
        "Bluetooth and Wi-Fi Direct, and speaks it again on the receiver. No tower, internet or server.")


# ----------------------------------------------------------------------------- slide 2
def slide2(s):
    content_slide(s, "iTantra: Send the Meaning, Not the Audio")
    heading(s, 0.5, 1.28, 12.3, "01", "Proposed Solution (Describe your Idea/Solution/Prototype)")
    text(s, 1.02, 1.66, 11.8, 0.3,
         "Detailed explanation: speech is recognised on the phone, sent as a few hundred bytes of signed text over a "
         "phone-to-phone mesh, and spoken again on the receiving phone.", size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

    # hero pipeline render, with labels under each stage
    iw = 9.6
    ih = iw / 3.322
    ix = (13.333 - iw) / 2
    iy = 1.98
    pic(s, os.path.join(IMG, "v2_pipeline_c.png"), ix, iy, w=iw)
    stages = [
        (0.09, "saff", "mic", "1  Speak + on-device STT", "VAD finds the pause, IndicConformer writes it"),
        (0.47, "emer", "mesh", "2  234 B packet hops the mesh", "BLE + Wi-Fi Direct relays, up to 7 hops, no tower"),
        (0.88, "navy", "speaker", "3  On-device TTS speaks it", "Voice note; ALERTs ring at max volume"),
    ]
    ly = iy + ih - 0.45
    for fx, g, ic, t, sub in stages:
        cx = ix + iw * fx
        w = 3.2
        x = min(max(cx - w / 2, 0.5), 12.83 - w)
        card(s, x, ly, w, 0.62, radius=0.2)
        badge(s, ic, g, x + 0.1, ly + 0.11, 0.44)
        text(s, x + 0.64, ly + 0.05, w - 0.7, 0.3, t, size=11, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.64, ly + 0.33, w - 0.7, 0.28, sub, size=8.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

    # bottom left: bandwidth comparison (how it addresses the problem)
    by = 5.14
    lx, lw = 0.5, 6.05
    card(s, lx, by, lw, 1.74, radius=0.06)
    text(s, lx + 0.2, by + 0.1, lw - 0.4, 0.3, [[("How it addresses the problem:  ", {"bold": True, "color": NAVY}),
                                               ("bytes to send one 3.5 s sentence", {"color": MUTED})]], size=11.5,
         anchor=MSO_ANCHOR.MIDDLE)
    bars = [("Raw voice, 64 kbps", 28000, "28,000 B", RGBColor(0xCB, 0xD5, 0xE1)),
            ("Opus codec, 16 kbps", 7000, "7,000 B", RGBColor(0x94, 0xA3, 0xB8)),
            ("iTantra packet", 234, "234 B", SAFF)]
    full = lw - 2.2 - 0.9
    for i, (lab, v, vs, col) in enumerate(bars):
        y = by + 0.45 + i * 0.32
        text(s, lx + 0.2, y, 1.9, 0.3, lab, size=10, color=SLATE, bold=(i == 2), anchor=MSO_ANCHOR.MIDDLE)
        bw = max(full * v / 28000, 0.06)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, lx + 2.15, y + 0.06, bw, 0.18, fill=col, radius=0.5)
        text(s, lx + 2.2 + bw + 0.05, y, 1.2, 0.3, vs, size=10, bold=True, color=NAVY if i < 2 else SAFF, anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, lx + 3.35, by + 1.09, 1.55, 0.3, grad=GRAD["saff"], radius=0.5, angle=0)
    text(s, lx + 3.35, by + 1.09, 1.55, 0.3, "120× smaller", size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    tags = [("offline", "No tower, internet or server"), ("hearing", "Heard by people who can't read"), ("phone", "Runs on 3 GB phones")]
    tx = lx + 0.2
    for ic, lab in tags:
        w = 0.054 * len(lab) + 0.48
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, tx, by + 1.4, w, 0.28, fill=T_EMER, radius=0.5)
        icon(s, ic, "g", tx + 0.08, by + 1.43, 0.22)
        text(s, tx + 0.33, by + 1.4, w - 0.36, 0.28, lab, size=9, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        tx += w + 0.08

    # bottom right: innovation & uniqueness
    rx, rw = 6.78, 6.05
    text(s, rx, by - 0.02, rw, 0.3, [[("Innovation and uniqueness", {"bold": True, "color": NAVY})]], size=12,
         anchor=MSO_ANCHOR.MIDDLE)
    feats = [
        ("wave", "saff", "Semantic radio", "Sends meaning + prosody, not the waveform"),
        ("alert", "saff", "Alert-first mesh", "Floods all hops; loud until acknowledged"),
        ("lock", "navy", "Tamper-proof", "Ed25519-signed; private messages encrypted"),
        ("offline", "emer", "100% on-device", "No cloud API; packs download once"),
    ]
    cw, ch = (rw - 0.14) / 2, 0.68
    for i, (ic, g, t, sub) in enumerate(feats):
        cx = rx + (i % 2) * (cw + 0.14)
        cy = by + 0.3 + (i // 2) * (ch + 0.06)
        card(s, cx, cy, cw, ch, radius=0.12)
        badge(s, ic, g, cx + 0.12, cy + 0.1, 0.48, glow=True)
        text(s, cx + 0.74, cy + 0.05, cw - 0.8, 0.3, t, size=12, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + 0.74, cy + 0.34, cw - 0.8, 0.3, sub, size=9.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "The bar compares bytes for one 3.5 second sentence: 64 kbps raw voice is 28,000 bytes, Opus 16 kbps is 7,000, "
        "our signed packet is 234 bytes (22 B header, text, 64 B Ed25519 signature).")


# ----------------------------------------------------------------------------- slide 3
def slide3(s):
    content_slide(s, "TECHNICAL APPROACH")
    lx, lw = 0.5, 7.55
    heading(s, lx, 1.28, lw, "02", "Technologies to be used")
    card(s, lx, 1.72, lw, 1.86, radius=0.06)
    rows = [
        ("navy", "APP & UI", ["Flutter (Dart)", "Kotlin (Android)", "Pigeon bridge", "Riverpod"]),
        ("saff", "ON-DEVICE AI", ["sherpa-onnx", "IndicConformer int8 (STT)", "MMS VITS (TTS)", "Silero VAD"]),
        ("emer", "MESH & SECURITY", ["BLE GATT mesh", "Wi-Fi Direct", "Ed25519", "X25519 + ChaCha20"]),
        ("sky", "MODEL TOOLING", ["Python", "ONNX int8", "Google Colab", "Hugging Face packs"]),
    ]
    for i, (g, cat, items) in enumerate(rows):
        y = 1.86 + i * 0.42
        shape(s, MSO_SHAPE.OVAL, lx + 0.2, y + 0.09, 0.14, 0.14, grad=GRAD[g], angle=45)
        text(s, lx + 0.42, y, 1.55, 0.32, cat, size=8.5, bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        x = lx + 1.95
        for it in items:
            x += chip(s, x, y + 0.01, it, fill=ICE, color=SLATE, line=HAIR, size=9.5) + 0.08

    heading(s, lx, 3.72, lw, "03", "Methodology and process for implementation (flow)")
    card(s, lx, 4.16, lw, 2.72, radius=0.05)

    def node(cx, cy, ic, g, t, sub):
        pulse(s, cx, cy, 0.34, color=GRAD[g][1], n=1, alpha=(30,), lw=0.75)
        badge(s, ic, g, cx - 0.26, cy - 0.26, 0.52)
        text(s, cx - 0.85, cy + 0.36, 1.7, 0.24, t, size=9.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        text(s, cx - 0.85, cy + 0.54, 1.7, 0.24, sub, size=8, color=MUTED, align=PP_ALIGN.CENTER)

    xs = [lx + 1.2 + i * 1.8 for i in range(4)]

    def vlabel(x, y, w, h, label, fill, color):
        cx, cy = x + w / 2, y + h / 2
        b = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - h / 2, cy - w / 2, h, w, fill=fill, radius=0.5)
        b.rotation = 270
        t = text(s, cx - h / 2, cy - w / 2, h, w, label, size=7.5, bold=True, color=color, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        t.rotation = 270

    vlabel(lx + 0.16, 4.3, 0.28, 1.1, "SENDER", T_SKY, NAVY2)
    sender = [("mic", "navy", "Mic 16 kHz", "echo + noise cancel"), ("wave", "navy", "Silero VAD", "pause = sentence end"),
              ("stt", "saff", "IndicConformer", "int8 CTC STT"), ("packet", "saff", "Packetiser", "sign · fragment · TTL 7")]
    for i, (ic, g, t, sub) in enumerate(sender):
        node(xs[i], 4.72, ic, g, t, sub)
        if i < 3:
            connector(s, xs[i] + 0.4, 4.72, xs[i + 1] - 0.4, 4.72)
    # mesh band
    my = 5.52
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, lx + 0.16, my, lw - 0.32, 0.3, grad=(NAVY2, NAVY), radius=0.5, angle=0)
    icon(s, "bt", "w", lx + 0.34, my + 0.04, 0.22)
    icon(s, "wifi", "w", lx + 0.6, my + 0.04, 0.22)
    text(s, lx + 0.95, my, lw - 1.2, 0.3,
         [[("MESH  ", {"bold": True, "color": WHITE}),
           ("verify → dedup → relay with jitter (≤7 hops) → store-and-forward outbox", {"color": RGBColor(0xDB, 0xE7, 0xF7)})]],
         size=9.5, anchor=MSO_ANCHOR.MIDDLE)
    vlabel(lx + 0.16, 5.9, 0.28, 0.92, "RECEIVER", T_EMER, EMER2)
    recv = [("shield", "emer", "Verify", "drop unsigned / replayed"), ("speaker", "emer", "MMS VITS TTS", "sender's language"),
            ("hearing", "sky", "Voice note", "text shown with audio"), ("alert", "saff", "ALERT mode", "max volume, repeat x3")]
    for i, (ic, g, t, sub) in enumerate(recv):
        node(xs[i], 6.14, ic, g, t, sub)
        if i < 3:
            connector(s, xs[i] + 0.4, 6.14, xs[i + 1] - 0.4, 6.14)

    # prototype: phone render with the app UI on screen
    heading(s, 8.35, 1.28, 4.5, "04", "Working prototype (Android)")
    ph_h = 5.05
    ph_w = ph_h * 1536 / 2752
    phx, phy = 8.5, 1.74
    pulse(s, phx + ph_w / 2, phy + ph_h / 2, 1.55, color=EMER, n=3, step=0.3, alpha=(18, 11, 6))
    sx0, sy0 = phx + ph_w * 0.196, phy + ph_h * 0.125
    sw, sh = ph_w * (0.795 - 0.196), ph_h * (0.876 - 0.125)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, sx0 - 0.04, sy0 - 0.06, sw + 0.08, sh + 0.12, fill=WHITE, radius=0.08)
    pic(s, os.path.join(IMG, "v2_phone_t.png"), phx, phy, ph_w, ph_h)
    text(s, sx0 + 0.08, sy0 + 0.1, sw - 0.1, 0.28, [[("Talk · ", {}), ("हिन्दी", {"font": "Nirmala UI"})]],
         size=11, bold=True, color=NAVY)
    shape(s, MSO_SHAPE.RECTANGLE, sx0, sy0 + 0.44, sw, 0.24, fill=RGBColor(0xEE, 0xF2, 0xEF))
    text(s, sx0 + 0.04, sy0 + 0.44, sw - 0.06, 0.24,
         [[("● ", {"color": EMER}), ("STT ready  ", {}), ("● ", {"color": EMER}), ("Voice  ", {}), ("● ", {"color": EMER}), ("Offline", {})]],
         size=6.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    bx, bw = sx0 + 0.36, sw - 0.44
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, sy0 + 0.86, bw, 0.64, fill=RGBColor(0xA7, 0xF3, 0xD0), radius=0.3)
    text(s, bx + 0.1, sy0 + 0.88, bw - 0.15, 0.34, "हलो हलो", size=14, color=RGBColor(0x06, 0x4E, 0x3B), font="Nirmala UI",
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, bx + 0.1, sy0 + 1.2, bw - 0.15, 0.24, "Hindi · STT 236 ms", size=7, color=RGBColor(0x04, 0x78, 0x57),
         anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, sx0 + 0.08, sy0 + sh - 1.02, sw - 0.16, 0.28, fill=WHITE, line=HAIR, radius=0.3)
    text(s, sx0 + 0.16, sy0 + sh - 1.02, sw - 0.3, 0.28, "Or type a message", size=7, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    badge(s, "mic", "emer", sx0 + sw - 0.66, sy0 + sh - 0.68, 0.56, pad=0.26, glow=True)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, sx0 + 0.08, sy0 + sh - 0.52, 0.78, 0.25, fill=RGBColor(0xDC, 0xFC, 0xE7), radius=0.5)
    text(s, sx0 + 0.08, sy0 + sh - 0.52, 0.78, 0.25, "PTT | Call", size=7, color=SLATE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)

    callouts = [("644 ms", "STT model load", NAVY, 2.35),
                ("236 ms", "Hindi sentence decode", SAFF, 3.4),
                ("654 ms", "TTS voice load", NAVY, 4.45),
                ("10 / 10", "languages verified", EMER2, 5.5)]
    cx0 = 11.42
    for big, lab, col, cy in callouts:
        card(s, cx0, cy - 0.38, 1.45, 0.86, radius=0.14)
        text(s, cx0, cy - 0.34, 1.45, 0.44, big, size=18, bold=True, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx0 + 0.05, cy + 0.08, 1.35, 0.36, lab, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        connector(s, phx + ph_w - 0.12, cy + 0.05, cx0 - 0.04, cy + 0.05, color=RGBColor(0x9A, 0xA9, 0xBB), lw=1.0,
                  dash=True, arrow=False)
    text(s, 8.35, 6.72, 4.5, 0.2, "App screen redrawn from the running build; timings measured on OnePlus CPH2717, Android 16.",
         size=7.5, color=MUTED)
    s.notes_slide.notes_text_frame.text = (
        "Everything runs on the phone: Flutter UI, Kotlin native layer, sherpa-onnx for STT, TTS and VAD. "
        "The prototype runs on a OnePlus today; timings are measured on that phone.")


# ----------------------------------------------------------------------------- slide 4
def slide4(s):
    content_slide(s, "FEASIBILITY AND VIABILITY")
    lx, lw = 0.5, 6.0
    heading(s, lx, 1.28, lw, "05", "Analysis of the feasibility of the idea")
    tiles = [("10 / 10", "languages running fully offline", EMER2, "check", "emer"),
             ("~0.3 GB", "per language: speech + voice packs", NAVY, "storage", "navy"),
             ("<1.1 GB", "peak RAM (est.) fits 3 GB phones", SAFF, "memory", "saff")]
    tw = (lw - 0.24) / 3
    for i, (big, lab, col, ic, g) in enumerate(tiles):
        x = lx + i * (tw + 0.12)
        card(s, x, 1.72, tw, 1.12, radius=0.12)
        badge(s, ic, g, x + 0.14, 1.84, 0.4)
        text(s, x + 0.62, 1.8, tw - 0.66, 0.48, big, size=20, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.14, 2.3, tw - 0.24, 0.5, lab, size=9, color=SLATE)

    card(s, lx, 2.98, lw, 3.86, radius=0.05)
    cd = CategoryChartData()
    langs = ["Kannada", "Marathi", "Tamil", "Bengali", "Gujarati", "Telugu", "Odia", "Hindi", "English", "Malayalam"]
    cer = [1.0, 1.2, 2.0, 3.2, 3.3, 3.3, 6.0, 8.5, 10.2, 13.6]
    cd.categories = langs
    cd.add_series("Character error %", cer)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(lx + 0.12), Inches(3.05), Inches(lw - 0.24), Inches(3.18), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Round-trip test: character error % per language (lower is better)"
    tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
    tr.font.size, tr.font.bold, tr.font.color.rgb, tr.font.name = Pt(10.5), True, NAVY, BODY
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '0.0', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(8.5), True, SLATE
    for i, v in enumerate(cer):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = EMER if v <= 5 else (NAVY2 if v <= 10 else SAFF)
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.color.rgb = Pt(8), MUTED
    ca.format.line.color.rgb = HAIR
    va = ch.value_axis
    va.visible = False
    va.maximum_scale = 16
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xF1, 0xF4, 0xF8)
    legend = [("≤5%", EMER), ("5–10%", NAVY2), (">10%", SAFF)]
    x = lx + 0.25
    for lab, col in legend:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 6.3, 0.22, 0.14, fill=col, radius=0.5)
        text(s, x + 0.27, 6.24, 0.7, 0.26, lab, size=8.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        x += 0.95
    text(s, lx + 3.05, 6.2, lw - 3.2, 0.6,
         "Our TTS speaks a flood warning, our STT transcribes it back (sherpa-onnx, CPU). "
         "STT runs 5–20× faster than real time.", size=8, color=MUTED)

    rx, rw = 6.75, 6.1
    heading(s, rx, 1.28, rw, "06", "Potential challenges and risks  →  Strategies to overcome them", size=13.5)
    risks = [
        ("route", "Short range per hop", "~60–100 m per Bluetooth hop",
         "7-hop relay, BLE Coded PHY where supported, store-and-forward by moving phones"),
        ("hearing", "Hard speech", "Noise, soft voices, Malayalam / Odia",
         "Gain control + tuned VAD, distress hotword list, text shown with voice, real-speech test set"),
        ("battery", "Low-end phones", "RAM, battery drain, background kills",
         "int8 models, one language in RAM, duty-cycled scanning, foreground service"),
        ("shield", "Fake alerts, misuse", "Spoofing; mesh apps flagged in India (Jul 2026)",
         "Ed25519-signed packets, verified NDMA / ISRO key, positioned as authenticated alert tool"),
        ("book", "Voice licence", "MMS-TTS is non-commercial",
         "Swap to CC-BY voices trained on SYSPIN / IndicVoices-R before deployment"),
    ]
    rh, gap = 0.95, 0.08
    for i, (ic, t, sub, fix) in enumerate(risks):
        y = 1.72 + i * (rh + gap)
        card(s, rx, y, rw, rh, radius=0.12)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, rx + 0.08, y + 0.08, 2.55, rh - 0.16, fill=T_SLATE, radius=0.14)
        icon(s, ic, "o", rx + 0.18, y + 0.17, 0.3)
        text(s, rx + 0.56, y + 0.12, 2.0, 0.3, t, size=11, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, rx + 0.2, y + 0.46, 2.36, 0.4, sub, size=8.5, color=MUTED)
        badge(s, "arrow", "saff", rx + 2.72, y + rh / 2 - 0.15, 0.3, pad=0.2)
        mx = rx + 3.1
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, mx, y + 0.08, rw - 3.18, rh - 0.16, fill=T_EMER, radius=0.14)
        icon(s, "check", "g", mx + 0.1, y + rh / 2 - 0.13, 0.26)
        text(s, mx + 0.44, y + 0.1, rw - 3.7, rh - 0.2, fix, size=9.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Feasibility is shown, not claimed: the app runs today and all 10 languages pass a round-trip check. "
        "The RAM figure is an estimate to confirm with dumpsys meminfo on a 3 GB phone.")


# ----------------------------------------------------------------------------- slide 5
def slide5(s):
    content_slide(s, "IMPACT AND BENEFITS")
    stats = [
        ("85%", "of Indians speak one of our 10 languages as mother tongue", "Census 2011", SAFF, "lang", "saff"),
        ("~26%", "of Indians aged 7+ cannot read: voice still reaches them", "Census 2011", NAVY2, "hearing", "navy"),
        ("0", "towers, internet, SIM cards or servers needed", "fully offline", EMER2, "offline", "emer"),
        ("₹0", "data cost per message, on phones people already own", "no data plan", SKY, "rupee", "sky"),
    ]
    sw = (12.33 - 0.36) / 4
    for i, (big, lab, src, col, ic, g) in enumerate(stats):
        x = 0.5 + i * (sw + 0.12)
        card(s, x, 1.3, sw, 1.2, radius=0.12)
        pulse(s, x + 0.55, 1.9, 0.3, color=col, n=2, step=0.14, alpha=(30, 15), lw=0.75)
        badge(s, ic, g, x + 0.33, 1.68, 0.44)
        text(s, x + 1.0, 1.36, sw - 1.05, 0.56, big, size=28, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.0, 1.9, sw - 1.1, 0.42, lab, size=9, color=SLATE)
        text(s, x + 1.0, 2.26, sw - 1.1, 0.2, src, size=7.5, color=MUTED, bold=True)

    heading(s, 0.5, 2.66, 5.9, "07", "Potential impact on the target audience")
    iw = 5.95
    pic(s, os.path.join(IMG, "v2_impact_c.png"), 0.45, 3.1, w=iw)
    labels = [("Relief camps", 0.72, 0.08), ("Flooded villages", 0.06, 0.2), ("NDRF / SDRF rescue teams", 0.18, 0.77)]
    ih = iw / 1.509
    for lab, fx, fy in labels:
        w = 0.068 * len(lab) + 0.3
        card(s, 0.45 + iw * fx, 3.1 + ih * fy, w, 0.3, radius=0.5)
        text(s, 0.45 + iw * fx, 3.1 + ih * fy, w, 0.3, lab, size=9, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)

    rx, rw = 6.75, 6.08
    heading(s, rx, 2.66, rw, "08", "Benefits of the solution")
    cards = [
        ("access", "saff", "Social", "Voice in 10 languages for non-literate, elderly and visually impaired people; no typing."),
        ("rupee", "emer", "Economic", "Runs on existing low-cost phones; zero infrastructure, data or SMS cost; open models."),
        ("eco", "emer", "Environmental", "No diesel towers or generators; low-power Bluetooth saves battery in long outages."),
        ("satellite", "navy", "Strategic (ISRO)", "Byte-sized frames suit NavIC messaging and LoRa links; data stays on Indian devices."),
    ]
    cw, chh = (rw - 0.14) / 2, 1.34
    for i, (ic, g, t, sub) in enumerate(cards):
        cx = rx + (i % 2) * (cw + 0.14)
        cy = 3.1 + (i // 2) * (chh + 0.12)
        card(s, cx, cy, cw, chh, radius=0.1)
        badge(s, ic, g, cx + 0.16, cy + 0.16, 0.52, glow=True)
        text(s, cx + 0.8, cy + 0.16, cw - 0.9, 0.52, t, size=13, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + 0.18, cy + 0.74, cw - 0.32, chh - 0.8, sub, size=9.5, color=SLATE)
    fy = 3.1 + 2 * (chh + 0.12)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, rx, fy, rw, 6.84 - fy, grad=(NAVY2, NAVY), radius=0.25, angle=0, shadow=True)
    icon(s, "rocket", "w", rx + 0.18, fy + (6.84 - fy - 0.32) / 2, 0.32)
    text(s, rx + 0.62, fy, rw - 0.75, 6.84 - fy,
         [[("NEXT  ", {"bold": True, "color": SAFF2}),
           ("\"heard by N phones\" receipts  •  SMS gateway on signal  •  voice-preserving prosody  •  10-phone field trial",
            {"color": WHITE})]],
         size=10, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Census 2011: Hindi, Bengali, Marathi, Telugu, Tamil, Gujarati, Kannada, Odia and Malayalam together are ~85% of "
        "mother-tongue speakers; literacy 74% means ~26% of people aged 7+ cannot read.")


# ----------------------------------------------------------------------------- slide 6
def slide6(s):
    content_slide(s, "RESEARCH AND REFERENCES")
    heading(s, 0.5, 1.28, 12.3, "09", "Details / Links of the reference and research work")
    cols = [
        ("layers", "navy", "Foundation models", [
            ("AI4Bharat IndicConformer ASR", "MIT", "github.com/AI4Bharat/IndicConformerASR"),
            ("sherpa-onnx offline runtime", "Apache-2.0", "github.com/k2-fsa/sherpa-onnx"),
            ("Meta MMS-TTS voices", "VITS", "huggingface.co/facebook/mms-tts"),
            ("Silero VAD", "MIT", "github.com/snakers4/silero-vad"),
            ("IndicConformer sherpa int8 exports", "HF", "huggingface.co/parismitaglobalsolutions/indicconformer-sherpa-onnx"),
        ]),
        ("science", "saff", "Academic literature", [
            ("IndicVoices: Indian speech dataset", "ACL 2024", "arxiv.org/abs/2403.01926"),
            ("STCTS: speech at 80 bps (text + prosody)", "arXiv 2025", "arxiv.org/abs/2512.00451"),
            ("IndicVoices-R: Indian TTS corpus", "NeurIPS 2024", "github.com/AI4Bharat/IndicVoices-R"),
            ("SYSPIN open Indian TTS dataset", "IISc", "vaani.iisc.ac.in/dataset/syspindataset"),
            ("BitChat mesh protocol whitepaper", "Open", "github.com/permissionlesstech/bitchat"),
        ]),
        ("public", "emer", "Field precedents & data", [
            ("FireChat mesh, Hong Kong 2014", "500k installs", "en.wikipedia.org/wiki/FireChat"),
            ("BitChat in Nepal 2025", "48,781 / day", "forbes.com/sites/digital-assets"),
            ("Bridgefy mesh adds E2E encryption", "TechCrunch 2020", "techcrunch.com/2020/11/02"),
            ("Census of India 2011", "Language, literacy", "censusindia.gov.in"),
            ("NDMA disaster management", "Govt. of India", "ndma.gov.in"),
        ]),
    ]
    links = {
        "forbes.com/sites/digital-assets": "https://www.forbes.com/sites/digital-assets/2025/09/11/jack-dorseys-bitchat-gains-traction-during-nepals-unrest/",
        "techcrunch.com/2020/11/02": "https://techcrunch.com/2020/11/02/bridgefy-launches-end-to-end-encrypted-messaging-for-the-app-used-during-protests-and-disasters/",
    }
    cw = (12.33 - 0.3) / 3
    for i, (ic, g, t, refs) in enumerate(cols):
        x = 0.5 + i * (cw + 0.15)
        card(s, x, 1.72, cw, 3.72, radius=0.05)
        badge(s, ic, g, x + 0.18, 1.86, 0.46, glow=True)
        text(s, x + 0.76, 1.86, cw - 0.9, 0.46, t, size=13.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        for j, (name, tag, url) in enumerate(refs):
            y = 2.46 + j * 0.59
            href = links.get(url, "https://" + url)
            shown = url if len(url) < 44 else url.split("/")[0] + "/.../" + url.rstrip("/").split("/")[-1]
            text(s, x + 0.2, y, cw - 1.45, 0.26, name, size=9.5, bold=True, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
            tw_ = 0.058 * len(tag) + 0.2
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + cw - tw_ - 0.18, y + 0.02, tw_, 0.22,
                  fill={"navy": T_SKY, "saff": T_SAFF, "emer": T_EMER}[g], radius=0.5)
            text(s, x + cw - tw_ - 0.18, y + 0.02, tw_, 0.22, tag, size=7.5, bold=True,
                 color={"navy": NAVY2, "saff": SAFF, "emer": EMER2}[g], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(s, x + 0.2, y + 0.26, cw - 0.4, 0.22, [[(shown, {"color": SKY, "link": href})]], size=8.5,
                 anchor=MSO_ANCHOR.MIDDLE)

    y0 = 5.58
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, y0, 12.33, 1.26, grad=(RGBColor(0xF5, 0xF9, 0xFF), RGBColor(0xFF, 0xF7, 0xF0)),
          radius=0.14, angle=0, shadow=True)
    pic(s, LOGO, 0.62, y0 + 0.1, 1.06, 1.06)
    text(s, 1.8, y0 + 0.12, 4.4, 0.34, "Built, verified and open source", size=14, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    bx = 1.8
    for lab in ["Android prototype", "20 language packs", "10/10 round-trip"]:
        bx += chip(s, bx, y0 + 0.56, lab, fill=WHITE, color=NAVY, line=HAIR, size=9.5, h=0.3, bold=True) + 0.1
    text(s, 1.8, y0 + 0.9, 4.6, 0.26, "Scan to see the code and download the offline model packs.", size=8.5, color=MUTED)
    qrs = [("App source code", "github.com/nileshpatil6/SIH-Hexabits", ["github.com/nileshpatil6/", "SIH-Hexabits"]),
           ("Offline model packs", "huggingface.co/datasets/Mr66/itantra-packs", ["huggingface.co/datasets/", "Mr66/itantra-packs"])]
    for k, (lab, url, lines) in enumerate(qrs):
        qx = 6.75 + k * 3.05
        card(s, qx, y0 + 0.1, 2.95, 1.06, radius=0.12)
        pic(s, qr_png("https://" + url), qx + 0.08, y0 + 0.15, 0.96, 0.96)
        text(s, qx + 1.12, y0 + 0.2, 1.8, 0.9,
             [[(lab, {"bold": True, "color": NAVY, "size": 10.5})]] +
             [[(ln, {"color": SKY, "size": 8.5, "link": "https://" + url})] for ln in lines])
    s.notes_slide.notes_text_frame.text = "All links are clickable in the PDF export."


def drop_slide(prs, index):
    sld = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


def main():
    prs = Presentation(os.path.join(HERE, "template.pptx"))
    drop_slide(prs, 6)
    for fn, sl in zip([slide1, slide2, slide3, slide4, slide5, slide6], prs.slides):
        fn(sl)
    out = os.path.join(HERE, "iTantra_SIH2026_Hexabits.pptx")
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
